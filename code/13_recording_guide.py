"""Recording guide for the team's own video: slide-by-slide content, spoken script, on-screen actions and all images.
Numbers are read from outputs/ so the guide always matches the current analysis.
Output: video/Video Recording Guide.docx  +  video/slide_kit/images/*.png (the same images, ready to insert)"""
import os, json, shutil
import pandas as pd
import docx
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from project_config import YEARS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
J = lambda *p: os.path.join(ROOT, *p)
KIT = J("video", "slide_kit", "images"); os.makedirs(KIT, exist_ok=True)
ml = json.load(open(J("outputs", "ml", "ml_summary.json")))
T = pd.read_csv(J("outputs", "stats", "T6_hypothesis_tests.csv")).set_index("id")
cr = pd.read_csv(J("outputs", "stats", "T5_holiday_vs_population_share.csv")).set_index("faith")
bt = pd.read_csv(J("outputs", "stats", "T2_by_tradition.csv")).set_index("tradition")["share_of_all_festival_holidays_pct"]
mon = pd.read_csv(J("outputs", "stats", "T3_month_by_tradition_holiday_days.csv"), index_col=0).sum(axis=1)
sp = pd.read_csv(J("data", "processed", "state_profile.csv")).set_index("state")
fm = pd.read_csv(J("data", "processed", "festival_master.csv"))
obs = pd.read_csv(J("data", "processed", "festival_observations.csv"))
eco = pd.read_csv(J("data", "processed", "economic_footfall_clean.csv"))
eng = json.load(open(J("blog", "engagement", "engagement_metrics.json")))
A, B, C = ml["A"], ml["B"]["metrics"], ml["C"]
top = sp.hol_days_total_festival.sort_values(ascending=False); low = top[::-1]
MN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
peak = [MN[i - 1] for i in mon.sort_values(ascending=False).index[:3]]
n_unsure = int((~fm.reach_identifiable).sum()); n_f = fm.festival.nunique()
rf, base = B["Random Forest"]["macro_f1_mean"], B["Majority baseline"]["macro_f1_mean"]
growth = T.loc["H9", "interpretation"].split("growth ")[1].split("%")[0].split(".")[0]
holi_err = C["holdout_2026"][1]["abs_pct_error"]
ne = next(v for k, v in A["detail_clusters"].items() if "Tribal" in k or "Christian" in k)
mus = next((v for k, v in A["detail_clusters"].items() if "Muslim" in k), [])

# ---- images: copy with slide-order names so they are easy to find ----
IMG = {"when": ("F02_month_timeline_by_tradition.png", "05_when_month_by_tradition.png"),
       "map": ("F05_map_festival_holidays.png", "06_where_state_map.png"),
       "heat": ("F04_state_month_heatmap.png", "06b_state_month_heatmap_optional.png"),
       "demo": (os.path.join("..", "..", "video", "slides", "_demography_wide.png"), "07_demography_three_panels.png"),
       "div": ("F06_map_holiday_diversity.png", "07b_diversity_map_optional.png"),
       "clus": ("M05_cluster_map.png", "08_state_groups_map.png"),
       "clf": ("M03_classifier_results.png", "09_classifier_results.png"),
       "trade": ("F08_festive_trade_trend.png", "10_festive_trade_trend.png"),
       "sector": ("F09_diwali_sector_split.png", "10b_diwali_sector_split.png"),
       "foot": ("F10_footfall_density.png", "11_footfall_density.png")}
PATH = {}
for k, (src, dst) in IMG.items():
    s = os.path.normpath(J("outputs", "figures", src)); d = os.path.join(KIT, dst); shutil.copyfile(s, d); PATH[k] = d

TEAM = ["Jainam Jain", "Vrushan Patil", "Dhanush Chowke", "Aditya Tambe", "Vivek Jaiswal"]
SLIDES = [
 dict(t="Title", sec=20, who=0,
      on=["Title: What India's Holiday Calendar Reveals", "Subtitle: Festivals, demography and the economics of celebration, measured from official data",
          "All five names, B.E. Sem VII ECS, VIT, Guide: Prof. Uma Jaishankar"],
      say=["Hello everyone. We are " + ", ".join(TEAM[:-1]) + " and " + TEAM[-1] + ".",
           "This is our Data Analytics and Visualization project. We studied Indian festivals using official data: when they happen, where they are recognised, whose festivals they are, and what they are worth economically.",
           "In the next few minutes we will take you through our dataset, our statistics, our machine learning models and our public blog."],
      do=["Everyone on camera for the first line if you are recording with video, then cut to the slide.", "Keep this slide short. Do not read the subtitle aloud."]),
 dict(t="Why we did this", sec=30, who=0,
      on=["Festival information is scattered and anecdotal", "We wanted evidence: when, where, whose, and what it is worth", "Our rule: zero synthetic data, every record has its source link"],
      say=["India celebrates hundreds of festivals. But if you ask which ones matter most, the answer usually depends on where the person grew up.",
           "Most information is scattered, and that makes biased comparisons easy. We did not want to rank festivals or favour any community.",
           "So we set one strict rule. No made-up data. Every single record in our dataset carries the link it came from."],
      do=["Bring the three points in one at a time as you say them.", "Stress the last line: this is the point your professor cares about."]),
 dict(t="Our dataset", sec=45, who=0,
      on=[f"{len(obs):,} festival-holiday records", f"{n_f} festivals, 9 tradition groups", "29 states and UTs (34 RBI offices)", f"{YEARS[0]} to {YEARS[-1]} (5 years)",
          f"{len(eco)} spending and crowd figures, each verified", "Sources: RBI, Census 2011, DoPT, Wikipedia and Government portals, CAIT and verified news"],
      say=[f"Our main source is the Reserve Bank of India. RBI publishes the official bank holidays for each of its thirty-four regional offices. We collected every holiday from {YEARS[0]} to {YEARS[-1]}.",
           f"That gave us {len(obs):,} records covering {n_f} festivals across twenty-nine states. They span Hindu, Muslim, Christian, Sikh, Buddhist, Jain, Parsi and tribal festivals.",
           "For population we used Census 2011, which is the latest religion census. Each festival's tradition comes from its Wikipedia page or a government portal.",
           f"We also collected {len(eco)} spending and crowd figures, and we opened every source page to confirm each number."],
      do=["Show the four big numbers as tiles. Big numbers read better than a paragraph.", "Optional: show a 5-second screen recording of the RBI holiday page, then the CSV file with its source-link column."]),
 dict(t="Keeping the count fair", sec=55, who=1,
      on=["Sundays: RBI lists no Sunday dates, so we average each festival only over the years it could be listed",
          "Shared dates: a festival a state is confirmed to observe counts in full",
          "Evidence used: RBI, Central Government compulsory-holiday list, each festival's home community",
          f"Still uncertain: {n_unsure} of {n_f} festivals (flagged)", "Not covered: festivals with no bank holiday, district-level holidays, UTs without an RBI office"],
      say=["Our professor asked us a fair question: how accurate is this? We found two things in the raw data that could quietly under-count a festival, and we corrected both.",
           "First, Sundays. Banks are closed on Sundays anyway, so RBI does not list them. If Christmas falls on a Sunday, it simply disappears that year. So we average each festival only over the years it could appear.",
           "Second, shared dates. RBI gives one combined label per date for the whole country, for example Ambedkar Jayanti, Vishu, Bihu and Baisakhi together. We did not split that day blindly. If a state is confirmed to observe a festival, it counts in full there.",
           f"We are also honest about the limits. {n_unsure} of the {n_f} festivals always share their date, so we flag them as uncertain. And festivals that have no bank holiday at all are not in this data.",
           "So what we measure is official recognition, not how many people celebrate."],
      do=["Use two boxes side by side: 'What we corrected' (green) and 'What the data cannot see' (amber).", "Speak slowly here. This slide answers the accuracy question directly."]),
 dict(t="When India celebrates", sec=40, who=1, img="when",
      on=[f"Peaks: {peak[0]}, {peak[1]}, {peak[2]}", "Not evenly spread (chi-square, p < 0.001)", f"Season does not depend on faith (p = {T.loc['H2', 'p_value']:.2f})"],
      say=[f"First finding: timing. Festival holidays are not spread evenly through the year. They peak in {peak[0]}, {peak[1]} and {peak[2]}.",
           "March and April carry Holi, Eid, Good Friday, Mahavir Jayanti and the regional new years. October is the Navratri, Durga Puja, Dussehra and Diwali season.",
           "A chi-square test confirms this is not chance. But here is the interesting part: a festival's season does not depend on its faith. Every tradition has festivals spread across the year."],
      do=["Point to the three tall bars as you name the months.", "Point to the colours: each colour is one tradition, and the legend is on the right."]),
 dict(t="Where India celebrates", sec=40, who=1, img="map", img2="heat",
      on=[f"Most: {top.index[0]} {top.iloc[0]:.1f}, {top.index[1]} {top.iloc[1]:.1f}, {top.index[2]} {top.iloc[2]:.1f}",
          f"Fewest: {low.index[0]} {low.iloc[0]:.1f}, {low.index[1]} {low.iloc[1]:.1f}, {low.index[2]} {low.iloc[2]:.1f}",
          f"National average: about {sp.hol_days_total_festival.mean():.0f} a year", f"No difference between the six regions (p = {T.loc['H6-hol_days', 'p_value']:.2f})"],
      say=[f"Second: geography. {top.index[0]}, {top.index[1]} and {top.index[2]} recognise the most festival holidays, around twenty-two a year.",
           f"{low.index[0]}, {low.index[1]} and {low.index[2]} have the fewest.",
           "But when we compared North, South, East, West, Central and the North-East, we found no significant difference. The variation is between individual states, not between regions."],
      do=["Darker blue means more festival holidays. Point to the darkest and lightest states.", "Optional second image: the state-by-month heatmap, to show each state's own festival season (for example Tamil Nadu in January)."]),
 dict(t="Do calendars mirror demography?", sec=55, who=2, img="demo",
      on=[f"Correlation {cr.spearman_rho.min():.2f} to {cr.spearman_rho.max():.2f} for all six faiths tested (all significant)",
          f"Hindu festivals: {bt['Hindu']:.0f}% of festival holidays vs 79.8% of population",
          f"Muslim {bt['Muslim']:.0f}% (pop. 14.2%), Christian {bt['Christian']:.0f}% (2.3%), Sikh {bt['Sikh']:.0f}% (1.7%)",
          "Reason: several minority festivals are recognised almost nationwide"],
      say=["This was the question we cared about most. Does a state's holiday calendar reflect who lives there?",
           f"The answer is yes. For every faith we tested, the share of holidays rises with that community's share of the population. The correlations run from {cr.spearman_rho.min():.2f} to {cr.spearman_rho.max():.2f}, and all are statistically significant. The Muslim one is the weakest.",
           f"But the calendar is also more plural than the population. Hindus are about eighty percent of India, but Hindu festivals are about {bt['Hindu']:.0f} percent of recognised festival holidays.",
           "Why? Because Christmas, Good Friday, both Eids, Buddha Purnima, Mahavir Jayanti and Guru Nanak Jayanti are recognised almost everywhere, whatever the local population. Hindu festivals are more numerous, but more regional."],
      do=["Each dot is one state. Trace the dashed trend line with the cursor: it goes up in all three panels.", "Pause after 'more plural than the population'. It is the headline finding."]),
 dict(t="Machine learning 1: grouping states", sec=45, who=2, img="clus",
      on=[f"K-Means on each state's holiday mix", f"Strongest split: {A['best_k']} groups (silhouette {A['silhouette']:.2f})",
          "Distinct group: " + ", ".join(ne), f"{A['detail_k']}-group view (silhouette {A['detail_silhouette']:.2f}) also separates " + ", ".join(mus),
          "The rest is a continuum, not regional blocks"],
      say=["Now the machine learning. Our first model groups states by the mix of their festival calendar, using K-Means clustering.",
           f"The strongest result is a split into two groups. {', '.join(ne[:-1])} and {ne[-1]} stand apart. Their calendars have far more Christian and tribal holidays.",
           f"A three-group view also separates {', '.join(mus[:-1])} and {mus[-1]}, where Muslim festivals are about a third of festival holidays." if mus else "",
           "Everything else is a continuum. The grouping does not follow the usual North, South, East, West blocks. Neighbouring states can have quite different calendars."],
      do=["Point to the north-eastern states first, then the three orange states, then sweep across the large blue area.", "Say 'silhouette score' only once. Do not explain the formula unless asked."]),
 dict(t="Machine learning 2: can a festival's footprint reveal its faith?", sec=45, who=3, img="clf",
      on=[f"Random Forest macro-F1: {rf:.2f}", f"Baseline (always guess the largest group): {base:.2f}", "Strongest signal: presence in the North-East",
          "Diwali and Holi get confused with the Eids and Christmas"],
      say=["Our second model asks a different question. If you only know a festival's date and which states observe it, can you tell which tradition it belongs to?",
           f"We trained a Random Forest and tested it with repeated cross-validation. It scores {rf:.2f}, against {base:.2f} for a simple baseline. So it learns something, but it is far from perfect.",
           "And that is the finding. The model confuses Diwali and Holi with the Eids and Christmas, because India's big festivals share the same nationwide footprint. In calendar terms, they look more alike than different."],
      do=["Left side is the confusion matrix, right side is feature importance. Point to the top bar (North-East presence).", "Be upfront that the score is modest. An honest modest result is better than an inflated one."]),
 dict(t="Machine learning 3 and economics: festive trade", sec=50, who=3, img="trade", img2="sector",
      on=["Diwali 2025: Rs 6.05 lakh crore (CAIT)", f"Growth: about {growth}% a year", f"2026 Holi estimate predicted within {holi_err:.0f}%", "Top Diwali sectors: grocery and FMCG 12%, jewellery 10%, electronics 8%"],
      say=["Third, the economics. The traders' body CAIT reports Diwali 2025 trade of six point zero five lakh crore rupees, up from one point two five lakh crore in 2021.",
           f"Our regression model estimates that reported festive trade is growing at about {growth} percent a year.",
           f"To test it, we trained the model only on data up to 2025 and asked it to predict 2026. It came within {holi_err:.0f} percent of CAIT's Holi estimate.",
           "These are survey estimates, and many are made before the festival. On the chart, filled dots are reported figures and open dots are projections."],
      do=["Trace the blue Diwali line upward.", "Optional second image: the Diwali sector chart. Name only the top three sectors."]),
 dict(t="Crowds, and a blind spot", sec=40, who=3, img="foot",
      on=["Maha Kumbh 2025: 66.21 crore visits in 45 days", "About 1.47 crore visits a day", "Blind spot: national spending figures exist mainly for Hindu festivals", "This shows what gets measured, not what matters"],
      say=["On crowds, Maha Kumbh 2025 recorded sixty-six crore visits in forty-five days. That is about one and a half crore a day, far denser than Oktoberfest or the Hajj.",
           "While collecting this, we noticed a blind spot. National spending estimates are published regularly for Hindu festivals, but rarely for Eid, Christmas, Gurpurab or tribal festivals.",
           "That does not mean those festivals matter less. It shows what gets measured. We think this gap is itself an important finding."],
      do=["Point to the orange bars (Indian festivals) and the grey bars (global comparisons).", "Slow down for the blind-spot point."]),
 dict(t="Our blog and the public response", sec=35, who=4,
      on=["Published on Medium (add the date)", "9 charts, every figure cited, secular framing", "External reactions and comments: [fill in] of 40", "Comments we replied to: [fill in]", "Code and data public on GitHub"],
      say=["We published our findings as a public article on Medium, with nine charts and a source for every figure. We kept the tone balanced. No festival is ranked above another.",
           "[Say your real numbers here.] So far the article has received ___ reactions and comments from people outside our college, and we have replied to ___ of them.",
           "All our code and data are public on GitHub, so anyone can check our work."],
      do=["Use a screenshot of the live Medium article, and a screenshot of the comments or reactions.", "Fill in the real numbers before recording. Do not guess them."]),
 dict(t="What the data shows", sec=35, who=4,
      on=[f"1. The calendar peaks in {peak[0]}, {peak[1]} and {peak[2]}", "2. Calendars follow demography, yet are more plural than the population",
          "3. Four north-eastern states are distinct; the rest is a continuum", "4. Festive trade is growing fast, but measured unevenly", "Thank you"],
      say=["To sum up, four things.",
           "One, India's official festival calendar peaks in spring and in October.",
           "Two, holiday calendars follow demography, yet they are more plural than the population.",
           "Three, four north-eastern states have a clearly distinct calendar, and the rest of India is a continuum.",
           "Four, festive trade is growing fast, but it is measured unevenly across traditions.",
           "Thank you. We are happy to take questions."],
      do=["Bring each point in as you say it.", "End on a clean thank-you slide with all five names and the GitHub and Medium links."]),
]
for s in SLIDES:
    s["say"] = [x for x in s["say"] if x]
    words = sum(len(x.split()) for x in s["say"])
    s["sec"] = int(round((words / 140 * 60 + 2) / 5) * 5)     # 140 words a minute, plus a pause; rounded to 5 s
total = sum(s["sec"] for s in SLIDES)

# ------------------------------------------------ document ------------------------------------------------
d = docx.Document()
sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size, color in [("Heading 1", 18, "0D366B"), ("Heading 2", 14, "0D366B"), ("Heading 3", 12, "256ABF")]:
    h = d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor.from_string(color)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
NAVY = "0D366B"
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tcPr.append(sh)
def borders(tb, color="BFC7D5"):
    pr = tb._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color); b.append(el)
    pr.append(b)
def bullets(cell_or_doc, items, style="List Bullet", size=11, italic=False):
    for it in items:
        p = cell_or_doc.add_paragraph(style=style); r = p.add_run(it); r.font.size = Pt(size); r.italic = italic
        p.paragraph_format.space_after = Pt(2)
def para(container, text, bold=False, size=11, color=None, italic=False, after=6):
    p = container.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after); return p

t = d.add_paragraph(); r = t.add_run("Video Recording Guide"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(NAVY)
para(d, "Socio-Economic & Demographic Analysis of Indian Festivals", bold=True, size=14, after=2)
para(d, "DAV mini project  |  " + ", ".join(TEAM) + "  |  B.E. Sem VII ECS, VIT", size=10.5, color="52514E", after=10)
para(d, f"This guide has everything needed to build your own slides and record the video: {len(SLIDES)} slides, what goes on each one, the words to say, "
        f"what to do on screen, and every chart. Total speaking time is about {total // 60} min {total % 60} s.", after=8)

d.add_heading("1. Before you record", level=1)
d.add_heading("Where the images are", level=3)
para(d, "Every chart in this guide is also saved as a full-size PNG, numbered in slide order, in this folder:", after=2)
para(d, r"D:\DAV Lab\project\video\slide_kit\images", bold=True, after=6)
para(d, "Insert them into PowerPoint from there. Do not screenshot them from this document: the originals are much sharper.", after=8)
d.add_heading("Slide design tips", level=3)
bullets(d, ["One idea per slide. Put the finding in the title, for example 'Festival holidays peak in April and October', not 'Timing'.",
            "Show at most three or four short points beside a chart. The words below are for you to say, not to paste onto the slide.",
            "Make key numbers big (for example 49% vs 79.8%). Big numbers are remembered; paragraphs are not.",
            "Use a 16:9 slide, a plain light background, and the same font throughout. Keep text at 24 pt or larger.",
            "Charts already have their own titles and source lines. Crop the chart title off if it repeats your slide title."])
d.add_heading("Recording tips", level=3)
bullets(d, ["Easiest tool: PowerPoint > Record > From Beginning. It records your voice per slide, and you can redo a single slide. Then Export > Create a Video > Full HD (1080p).",
            "Use earphones with a mic, in a quiet room. Do one test slide and listen back before recording everything.",
            "Speak a little slower than feels natural. Pause for a second when you change slides.",
            "Do not read word for word in a flat voice. Read the script twice, then say it in your own way. Small changes in wording are fine; the numbers must stay exact.",
            "If more than one person speaks, record each person's slides in one sitting so the sound stays consistent.",
            f"The full script runs about {total // 60} min {total % 60} s at a normal pace. If you need it shorter, drop the optional images and trim slides 3, 10 and 11 first."])
d.add_heading("Suggested speaker split", level=3)
tb = d.add_table(rows=1, cols=3); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(tb)
for c, h in zip(tb.rows[0].cells, ["Speaker", "Slides", "Covers"]):
    c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.color.rgb = RGBColor(255, 255, 255); shade(c, NAVY)
split = {}
for i, s in enumerate(SLIDES, 1): split.setdefault(s["who"], []).append(i)
cover = ["Introduction, motivation, dataset", "Data accuracy, timing, geography", "Demography, state grouping model", "Classifier, trade model, crowds", "Blog response, conclusion"]
for w, idx in split.items():
    row = tb.add_row().cells
    row[0].text = TEAM[w]; row[1].text = f"{idx[0]} to {idx[-1]}" if len(idx) > 1 else str(idx[0]); row[2].text = cover[w]
for row in tb.rows:
    for c, wd in zip(row.cells, [Cm(4.5), Cm(3), Cm(9.5)]): c.width = wd
para(d, "This is only a suggestion. Swap names freely, or have one person narrate everything.", size=10, italic=True, color="52514E")

d.add_page_break()
d.add_heading("2. Slide by slide", level=1)
para(d, "For each slide: what to put on it, what to say, what to do on screen, and the image to use.", after=4)
for i, s in enumerate(SLIDES, 1):
    d.add_page_break() if i > 1 else None
    d.add_heading(f"Slide {i}: {s['t']}", level=2)
    para(d, f"About {s['sec']} seconds   |   Speaker: {TEAM[s['who']]}", size=10, color="52514E", after=4)
    tb = d.add_table(rows=0, cols=2); borders(tb)
    def row(label, fill):
        cells = tb.add_row().cells; cells[0].width = Cm(3.2); cells[1].width = Cm(13.8)
        cells[0].text = ""; rr = cells[0].paragraphs[0].add_run(label); rr.bold = True; rr.font.size = Pt(10.5); shade(cells[0], fill)
        cells[1].text = ""; return cells[1]
    c = row("Put on the slide", "EEF3FA"); first = True
    for it in s["on"]:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.style = d.styles["List Bullet"]; p.add_run(it).font.size = Pt(10.5); p.paragraph_format.space_after = Pt(1)
    c = row("Say", "E6F5EC"); first = True
    for it in s["say"]:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.add_run(it).font.size = Pt(11); p.paragraph_format.space_after = Pt(5)
    c = row("Do", "FDF1D8"); first = True
    for it in s["do"]:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.style = d.styles["List Bullet"]; p.add_run(it).font.size = Pt(10.5); p.paragraph_format.space_after = Pt(1)
    for key, lab in [("img", "Image"), ("img2", "Optional second image")]:
        if s.get(key):
            pth = PATH[s[key]]
            para(d, f"{lab}: {os.path.basename(pth)}", bold=True, size=10, color="256ABF", after=2).paragraph_format.space_before = Pt(8)
            from PIL import Image as _I
            w_, h_ = _I.open(pth).size; width = min(16.5, 9.0 * w_ / h_) if key == "img" else min(11.0, 7.0 * w_ / h_)
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(pth, width=Cm(width))
    if not s.get("img"):
        note = ("No chart on this slide. Use your own screenshots of the live Medium article and its reactions." if "blog" in s["t"].lower()
                else "No chart on this slide. Use text, big numbers or simple boxes.")
        para(d, note, size=10, italic=True, color="52514E").paragraph_format.space_before = Pt(8)

d.add_page_break()
d.add_heading("3. Numbers to get exactly right", level=1)
para(d, "Keep this page open while recording. These are the figures most likely to be misquoted.", after=4)
tb = d.add_table(rows=1, cols=2); borders(tb)
for c, h in zip(tb.rows[0].cells, ["Fact", "Value"]):
    c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.color.rgb = RGBColor(255, 255, 255); shade(c, NAVY)
facts = [("Records in the main dataset", f"{len(obs):,}"), ("Festivals / tradition groups", f"{n_f} / 9"), ("States and UTs / RBI offices", "29 / 34"),
         ("Years covered", f"{YEARS[0]} to {YEARS[-1]}"), ("Verified spending and crowd figures", f"{len(eco)}"),
         ("Festivals whose state coverage is uncertain", f"{n_unsure} of {n_f}"), ("Busiest months", ", ".join(peak)),
         ("Recognised festival holidays per year (national average)", f"about {sp.hol_days_total_festival.mean():.0f}"),
         ("Most festival holidays", f"{top.index[0]} {top.iloc[0]:.1f}, {top.index[1]} {top.iloc[1]:.1f}, {top.index[2]} {top.iloc[2]:.1f}"),
         ("Fewest festival holidays", f"{low.index[0]} {low.iloc[0]:.1f}, {low.index[1]} {low.iloc[1]:.1f}, {low.index[2]} {low.iloc[2]:.1f}"),
         ("Holiday share vs population share (Spearman)", f"{cr.spearman_rho.min():.2f} to {cr.spearman_rho.max():.2f}, all significant"),
         ("Hindu share of festival holidays vs population", f"{bt['Hindu']:.0f}% vs 79.8%"),
         ("Muslim / Christian / Sikh share of festival holidays", f"{bt['Muslim']:.0f}% / {bt['Christian']:.0f}% / {bt['Sikh']:.0f}%"),
         ("State grouping: best split", f"{A['best_k']} groups, silhouette {A['silhouette']:.2f}"),
         ("Distinct group of states", ", ".join(ne)),
         ("Classifier: Random Forest vs baseline (macro-F1)", f"{rf:.2f} vs {base:.2f}"),
         ("Festive trade growth", f"about {growth}% a year"), ("Diwali 2025 trade (CAIT)", "Rs 6.05 lakh crore"),
         ("2026 Holi estimate: model error", f"{holi_err:.0f}%"), ("Maha Kumbh 2025", "66.21 crore visits in 45 days"),
         ("Blog engagement", f"{eng['external_interactions']} of 40 logged so far: use your real, current numbers")]
for a_, b_ in facts:
    cells = tb.add_row().cells; cells[0].text = a_; cells[1].text = b_
for row_ in tb.rows:
    for c, wd in zip(row_.cells, [Cm(8.5), Cm(8.5)]):
        c.width = wd
        for p in c.paragraphs:
            for r_ in p.runs: r_.font.size = Pt(10.5)

d.add_heading("4. If you are asked", level=1)
qa = [("Why bank holidays?", "They are the only official, state-wise record that covers every faith in one format, for several years. They measure official recognition, not how many people celebrate."),
      ("Is Census 2011 too old?", "It is the latest religion census India has published. We say so as a limitation."),
      ("Are holidays declared later included?", "Yes. RBI updates its list, and our data includes those dates, for example election days and heavy-rain closures."),
      ("What about festivals on Sundays?", "RBI does not list Sundays, so we average each festival only over the years it could be listed."),
      ("Are the CAIT trade figures reliable?", "They are trader-survey estimates, and many are made before the festival. We mark each one as reported or projected."),
      ("Why is the classifier score low?", "There are only about a hundred festivals, and big festivals of different faiths look alike on the calendar. We report honest cross-validated results.")]
for q, a_ in qa:
    p = d.add_paragraph(); r1 = p.add_run(q + "  "); r1.bold = True; p.add_run(a_); p.paragraph_format.space_after = Pt(6)

out = J("video", "Video Recording Guide.docx"); d.save(out)
print("guide ->", out, "| slides:", len(SLIDES), "| speaking time s:", total, "| images:", len(os.listdir(KIT)))
